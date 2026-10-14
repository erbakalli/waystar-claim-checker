from dataclasses import replace
from datetime import date
from decimal import Decimal
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
for folder in (ROOT / "src", ROOT.parent / "Kantime-engine-main" / "src"):
    sys.path.insert(0, str(folder))

from kantime_engine import RowRecord, RowResult, RowDisposition, PortalAutomationError, RecoveryAction, RetryScope
from kantime_engine.runner import AutomationRunner
from kantime_engine.cancellation import CancellationToken
from kantime_engine.paths import ApplicationPaths
from kantime_engine.workbook import WorkbookSession
from waystar_claim_checker.application import Automation, claim_output, create_application
from waystar_claim_checker.claims import Claim, money, parse_table, select_claim
from waystar_claim_checker.pages import load_profile
from openpyxl import Workbook, load_workbook

EXPECTED_CLAIM_HEADERS = (
    "Sr #", "PatientName", "Office", "Contract", "VisitDate", "Visit/Supply/Expense",
    "InvoiceNumber", "Invoice Date", "Units", "Amount", "PaidAmount", "Adj.",
    "Balance", "Pay. Status", "Current Reason for Non - Payment", "TRN",
    "Note(s)/ Follow Up", "AltPatientID", "Representative", "DOB", "StartDate",
    "EndDate", "Claim ID", "Billed Amount", "Paid Amount", "Claim Status",
    "Denial Reason", "Last Action Taken", "Finalized Date", "Check Number",
    "Check Date", "Recheck", "AutomationStatus", "LastError",
)


def input_values(**updates):
    values = dict(invoice_no="INV-01", insurance_id="TEST-MBI-001", icn="",
                  service_from=date(2026, 5, 1), service_to=date(2026, 5, 31),
                  invoice_amount=100, comments="Preserve this existing note")
    values.update(updates)
    return values


def claim(**updates):
    return replace(Claim("TEST-MBI-001", "TEST-ICN-001", date(2026, 5, 1),
                         date(2026, 5, 31), Decimal("100.00"), Decimal("80.00"),
                         "Paid", "0000123", date(2026, 6, 10)), **updates)


class MatchingTests(unittest.TestCase):
    def test_input_invoice_amount_overrides_only_output_amount(self):
        automation = Automation()
        for amount in (Decimal("250.75"), 0, "$250.75", None, "", "   "):
            with self.subTest(amount=amount), patch("waystar_claim_checker.application.ClaimsPage") as page:
                page.return_value.search.return_value = [claim(invoice="INV-01")]
                result = automation.process_row(
                    RowRecord(0, "INV-01", input_values(invoice_amount=amount)),
                    SimpleNamespace(cancellation=CancellationToken()))
                self.assertEqual(result.disposition, RowDisposition.RESOLVED)
                output = result.values["claim_rows"][0]
                expected = Decimal("100.00") if amount is None or not str(amount).strip() else amount
                self.assertEqual(output["amount"], expected)
                self.assertEqual(result.values["amount"], expected)
                self.assertEqual(output["total_charges"], Decimal("100.00"))
                self.assertEqual(output["balance"], Decimal("20.00"))

    def test_legacy_results_are_reprocessed_after_search_validation_changes(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            source = Path(temp) / "legacy.xlsx"
            book = Workbook()
            book.active.append(["Invoice No"])
            book.active.append(["0001"])
            book.active.append(["0002"])
            book.save(source)
            book.close()
            app = create_application()
            paths = ApplicationPaths("legacy-search", "legacy-search", Path(temp) / "state")
            def session(spec):
                return WorkbookSession(spec, app.automation.spec.variants[0],
                                       plugin_id=app.spec.id, app_paths=paths, output_policy=app.spec.output)
            previous = session(replace(app.automation.spec.workbook, processing_revision=""))
            try:
                previous.prepare(str(source))
                batch, _ = previous.begin_batch(1)
                row = previous.claim_next(batch)
                previous.commit(row, RowResult(RowDisposition.RESOLVED,
                                values={"claim_status": "Not Found", "claim_rows": []}))
                previous.finish_batch()
            finally:
                previous.close()
            current = session(app.automation.spec.workbook)
            try:
                choose_resume = Mock(return_value=True)
                prepared = current.prepare(str(source), choose_resume)
                self.assertFalse(prepared.resumed)
                choose_resume.assert_not_called()
                batch, count = current.begin_batch(10)
                self.assertEqual(count, 2)
                self.assertEqual(current.claim_next(batch).key, "0001")
            finally:
                current.close()

    def test_workbook_with_only_invoice_number_can_search_and_export(self):
        for header in ("Invoice No", "Invoice Number", "InvoiceNumber"):
            with self.subTest(header=header), tempfile.TemporaryDirectory(dir=ROOT) as temp:
                source = Path(temp) / "invoice-only.xlsx"
                book = Workbook()
                book.active.append([header])
                book.active.append(["000123"])
                book.save(source)
                book.close()
                original = source.read_bytes()
                app = create_application()
                paths = ApplicationPaths("invoice-only", "invoice-only", Path(temp) / "state")
                session = WorkbookSession(app.automation.spec.workbook, app.automation.spec.variants[0],
                                          plugin_id=app.spec.id, app_paths=paths, output_policy=app.spec.output)
                try:
                    session.prepare(str(source))
                    batch, _ = session.begin_batch(10)
                    record = session.claim_next(batch)
                    with patch("waystar_claim_checker.application.ClaimsPage") as page:
                        page.return_value.search.return_value = [claim(invoice="000123", dcn="000456")]
                        result = app.automation.process_row(record, SimpleNamespace(cancellation=CancellationToken()))
                        page.return_value.search.assert_called_once_with("000123")
                    session.commit(record, result)
                    session.finish_batch()
                    report = load_workbook(session.export_report(), read_only=True, data_only=True)
                    try:
                        rows = list(report["Claims"].iter_rows(values_only=True))
                        self.assertEqual(rows[1][rows[0].index("InvoiceNumber")], "000123")
                        self.assertEqual(rows[1][rows[0].index("Claim ID")], "000456")
                        self.assertEqual(rows[1][rows[0].index("Amount")], 100)
                        self.assertIn("Comments", next(report["Results"].iter_rows(values_only=True)))
                    finally:
                        report.close()
                    self.assertEqual(source.read_bytes(), original)
                finally:
                    session.close()

    def test_search_read_error_is_retried_before_resolving_invoice(self):
        automation = Automation()
        token = CancellationToken()
        context = SimpleNamespace(cancellation=token, logger=Mock())
        runner = AutomationRunner(Mock(), token)
        error = PortalAutomationError("Transient results", reason="unrecognized_results",
                                     recovery=RecoveryAction.BACKOFF, retry_scope=RetryScope.CURRENT_ROW)
        with patch("waystar_claim_checker.application.ClaimsPage") as page, patch.object(runner, "_recover"):
            page.return_value.search.side_effect = [error, [claim(invoice="INV-01")]]
            result = runner._process_with_retry(automation, context, Mock(),
                                                RowRecord(0, "INV-01", input_values()), 3)
            self.assertEqual(page.return_value.search.call_count, 2)
        self.assertEqual(result.disposition, RowDisposition.RESOLVED)
        self.assertEqual(len(result.values["claim_rows"]), 1)

    def test_persistent_read_error_is_deferred_after_three_attempts(self):
        automation = Automation()
        token = CancellationToken()
        context = SimpleNamespace(cancellation=token, logger=Mock())
        runner = AutomationRunner(Mock(), token)
        workbook = Mock()
        workbook.blank_output_values.return_value = {}
        error = PortalAutomationError("Unreadable results", reason="unrecognized_results",
                                     recovery=RecoveryAction.BACKOFF, retry_scope=RetryScope.CURRENT_ROW)
        with patch("waystar_claim_checker.application.ClaimsPage") as page, patch.object(runner, "_recover"):
            page.return_value.search.side_effect = error
            result = runner._process_with_retry(automation, context, workbook,
                                                RowRecord(0, "INV-01", input_values()), 3)
            self.assertEqual(page.return_value.search.call_count, 3)
        self.assertEqual(result.disposition, RowDisposition.RETRY_LATER)
        self.assertEqual(result.comment, "Unreadable results")
        self.assertNotIn("claim_rows", result.values)

    def test_exact_period_and_amount(self):
        self.assertEqual(select_claim(input_values(), [claim()])[0], "matched")
        for change in ({"billed": Decimal("101")}, {"beneficiary_id": "OTHER"},
                       {"to_date": date(2026, 5, 30)}):
            self.assertEqual(select_claim(input_values(), [claim(**change)])[0], "not_found")

    def test_icn_requires_correct_beneficiary_and_never_falls_back(self):
        self.assertEqual(select_claim(input_values(icn="TEST-ICN-001"), [claim()])[0], "matched")
        self.assertEqual(select_claim(input_values(icn="OTHER"), [claim()])[0], "not_found")
        self.assertEqual(select_claim(input_values(icn="TEST-ICN-001"), [claim(beneficiary_id="OTHER")])[0], "not_found")

    def test_ambiguity_and_invoice_disambiguation(self):
        self.assertEqual(select_claim(input_values(), [claim(), claim()])[0], "ambiguous")
        self.assertEqual(select_claim(input_values(), [claim(invoice="INV-01"), claim(invoice="OTHER")])[0], "matched")

    def test_missing_and_invalid_inputs(self):
        self.assertEqual(select_claim(input_values(insurance_id=""), [claim()])[0], "missing_input")
        self.assertEqual(select_claim(input_values(invoice_amount=None), [claim()])[0], "missing_input")
        self.assertEqual(select_claim(input_values(service_to=date(2026, 4, 1)), [claim()])[0], "invalid_input")

    def test_money_preserves_missing_zero_and_negative(self):
        self.assertIsNone(money(""))
        self.assertEqual(money("$0.00"), Decimal("0.00"))
        self.assertEqual(money("($1,234.56)"), Decimal("-1234.56"))
        with self.assertRaises(ValueError):
            money("NaN")

    def test_table_parser_and_missing_payments(self):
        headers = ["MBI", "ICN", "From Date", "To Date", "Billed Amount", "Status"]
        parsed = parse_table(headers, [["TEST-MBI-001", "TEST-ICN-001", "05/01/2026", "05/31/2026", "$100.00", "Pending"]])
        self.assertEqual(select_claim(input_values(), parsed)[0], "matched")
        self.assertIsNone(parsed[0].paid)
        with self.assertRaises(ValueError):
            parse_table(["Follow-up On", "Patient Name"], [["", ""]])
        with self.assertRaises(ValueError):
            parse_table(headers, [["too short"]])

    def test_unverified_profile_rejected(self):
        with self.assertRaises(ValueError):
            load_profile(ROOT / "tests" / "portal-profile.fixture.json")

    def test_verified_profile_requires_pcn_selection_fields(self):
        profile = json.loads((ROOT / "tests" / "portal-profile.fixture.json").read_text())
        profile["verified"] = True
        profile.update(ready="#completed", no_results="#empty", submit="#submit")
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            path = Path(temp) / "profile.json"
            path.write_text(json.dumps(profile))
            self.assertEqual(load_profile(path)["search_input"], "#search-text")
            for key in ("initial_search_input", "search_mode_toggle", "pcn_option"):
                incomplete = dict(profile)
                del incomplete[key]
                path.write_text(json.dumps(incomplete))
                with self.assertRaises(ValueError):
                    load_profile(path)

    def test_placeholder_selectors_stay_disabled_even_if_verified_is_true(self):
        profile = json.loads((ROOT / "tests" / "portal-profile.fixture.json").read_text())
        profile["verified"] = True
        profile.update(ready="REPLACE_WITH_COMPLETED_SEARCH_SELECTOR",
                       no_results="REPLACE_WITH_NO_RESULTS_SELECTOR")
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            path = Path(temp) / "profile.json"
            path.write_text(json.dumps(profile))
            with self.assertRaises(ValueError) as caught:
                load_profile(path)
            self.assertIn("completed-search indicator", str(caught.exception))
            self.assertIn("no-results message", str(caught.exception))



    def test_deferred_search_keeps_original_comments(self):
        automation = Automation()
        error = PortalAutomationError("Search has more pages.", retry_scope=RetryScope.NEXT_RUN)
        with patch("waystar_claim_checker.application.ClaimsPage") as page:
            page.return_value.search.side_effect = error
            result = automation.process_row(RowRecord(0, "INV-01", input_values()),
                                            SimpleNamespace(cancellation=CancellationToken()))
        self.assertIsNone(result.comment)
        self.assertEqual(result.disposition, RowDisposition.RETRY_LATER)
        self.assertEqual(result.values["checker_comments"], str(error))

    def test_live_search_uses_invoice_no_as_pcn(self):
        automation = Automation()
        with patch("waystar_claim_checker.application.ClaimsPage") as page:
            page.return_value.search.return_value = [claim(invoice="INV-01")]
            result = automation.process_row(RowRecord(0, "INV-01", input_values()),
                                            SimpleNamespace(cancellation=CancellationToken()))
        page.return_value.search.assert_called_once_with("INV-01")
        self.assertEqual(result.values["claim_status"], "Found (1)")

    def test_waystar_headers_dcn_and_tooltip_description(self):
        headers = ["Status", "Claim Status", "Patient Full Name", "Beneficiary ID", "PCN",
                   "Stmt Start", "Stmt End", "Total Charges", "Primary Reason", "Paid Date", "DCN",
                   "Received Date", "Provider Reimb", "Paid Amount"]
        row = ["NEW", "Rejected", "EXAMPLE, PATIENT", "TEST-MBI-001", "INV-01", "05/01/26",
               "05/31/26", "$100.00", "12345", "06/10/26", "0000000123AZR",
               "06/01/26", "$75.00", "$90.00"]
        parsed = parse_table(headers, [row], primary_reasons=["Example rejection description."])
        self.assertEqual(parsed[0].status, "Rejected")
        self.assertEqual(parsed[0].primary_reason, "Example rejection description.")
        self.assertEqual(parsed[0].dcn, "0000000123AZR")
        self.assertEqual(parsed[0].paid_date, date(2026, 6, 10))
        self.assertEqual(parsed[0].claim_id, "")
        output = claim_output(parsed[0])
        self.assertEqual(output["pcn"], "INV-01")
        self.assertEqual(output["invoice_date"], date(2026, 6, 1))
        self.assertEqual(output["amount"], Decimal("100.00"))
        self.assertEqual(output["source_paid_amount"], Decimal("75.00"))
        self.assertEqual(output["paid_amount"], Decimal("75.00"))
        self.assertEqual(output["adjustment"], Decimal("25.00"))
        self.assertEqual(output["balance"], Decimal("25.00"))
        self.assertEqual(output["payment_status"], "Rejected")
        self.assertEqual(output["trn"], "0000000123AZR")
        self.assertEqual(output["alt_patient_id"], "TEST-MBI-001")
        self.assertEqual(output["claim_id"], "0000000123AZR")
        self.assertEqual(output["denial_reason"], "Example rejection description.")
        self.assertEqual(output["finalized_date"], date(2026, 6, 10))
        with self.assertRaises(ValueError):
            parse_table(headers, [row])
        with self.assertRaises(ValueError):
            parse_table(headers, [row], primary_reasons=["12345"])

    def test_all_claim_rows_survive_saved_state_and_export(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            source = Path(temp) / "synthetic.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Invoice No", "Insurance ID", "Statement Period Start Date",
                          "Statement Period End Date", "Invoice Amount", "Comments"])
            sheet.append(["INV-01", "", None, None, 250.75, "Keep original"])
            sheet.append(["INV-02", "", None, None, None, "No portal claims"])
            sheet.append(["INV-03", "", None, None, None, "Pending invoice"])
            book.save(source)
            book.close()
            app = create_application()
            app.automation.office = "Arizona Home Care"
            paths = ApplicationPaths("waystar-all-test", "waystar-all-test", Path(temp) / "state")
            def session():
                return WorkbookSession(app.automation.spec.workbook, app.automation.spec.variants[0],
                                       plugin_id=app.spec.id, app_paths=paths, output_policy=app.spec.output)
            current = session()
            try:
                current.prepare(str(source))
                batch, count = current.begin_batch(10)
                record = current.claim_next(batch)
                long_reason = "Synthetic explanation that must stay inside the reason cell. " * 12
                returned = [claim(invoice="INV-01", patient_full_name="EXAMPLE, PATIENT",
                                  primary_reason=long_reason, dcn="000123", paid_date=date(2026, 6, 10),
                                  received_date=date(2026, 6, 1)),
                            claim(invoice="INV-01", patient_full_name="EXAMPLE, PATIENT",
                                  primary_reason="Different reason", dcn="000456", status="Suspense",
                                  paid=None, check_date=None, received_date=date(2026, 6, 1))]
                with patch("waystar_claim_checker.application.ClaimsPage") as page:
                    page.return_value.search.return_value = returned
                    result = app.automation.process_row(record, SimpleNamespace(cancellation=CancellationToken()))
                self.assertEqual(len(result.values["claim_rows"]), 2)
                current.commit(record, result)
                record = current.claim_next(batch)
                with patch("waystar_claim_checker.application.ClaimsPage") as page:
                    page.return_value.search.return_value = []
                    result = app.automation.process_row(record, SimpleNamespace(cancellation=CancellationToken()))
                current.commit(record, result)
                current.finish_batch()
                run_id = current.run.id
            finally:
                current.close()
            restored = session()
            try:
                restored.run = restored.store.get_run(run_id)
                report = load_workbook(restored.export_report(), data_only=True)
                try:
                    rows = list(report["Claims"].iter_rows(values_only=True))
                    self.assertEqual(rows[0], EXPECTED_CLAIM_HEADERS)
                    headers = rows[0]
                    self.assertEqual(len(rows), 4)
                    self.assertEqual([row[headers.index("Office")] for row in rows[1:]],
                                     ["Arizona Home Care", "Arizona Home Care", None])
                    self.assertEqual([row[headers.index("Contract")] for row in rows[1:]], ["NGS", "NGS", None])
                    self.assertEqual(rows[1][headers.index("Sr #")], 1)
                    self.assertEqual(rows[1][headers.index("Current Reason for Non - Payment")],
                                     f"1. {long_reason}\n2. Different reason")
                    self.assertEqual(rows[1][headers.index("Claim Status")], "1. Paid\n2. Suspense")
                    self.assertEqual(rows[1][headers.index("Check Date")], "1. 06/10/2026\n2. ")
                    self.assertEqual(rows[1][headers.index("StartDate")].date(), date(2026, 5, 1))
                    self.assertEqual(rows[1][headers.index("EndDate")].date(), date(2026, 5, 31))
                    self.assertEqual(rows[1][headers.index("AltPatientID")], "TEST-MBI-001")
                    self.assertEqual(rows[1][headers.index("Denial Reason")], f"1. {long_reason}\n2. Different reason")
                    self.assertEqual(rows[1][headers.index("Finalized Date")], "1. 06/10/2026\n2. ")
                    self.assertEqual(rows[1][headers.index("VisitDate")], "05/01/2026 - 05/31/2026")
                    self.assertEqual(rows[1][headers.index("InvoiceNumber")], "INV-01")
                    self.assertEqual(rows[1][headers.index("Invoice Date")].date(), date(2026, 6, 1))
                    self.assertEqual(rows[1][headers.index("Amount")], 250.75)
                    self.assertEqual(rows[1][headers.index("PaidAmount")], "1. $80.00\n2. ")
                    self.assertEqual(rows[1][headers.index("Adj.")], "1. $20.00\n2. ")
                    self.assertEqual(rows[1][headers.index("Balance")], "1. $20.00\n2. ")
                    self.assertEqual(rows[1][headers.index("Pay. Status")], "1. Paid\n2. Suspense")
                    self.assertEqual(rows[1][headers.index("TRN")], "1. 000123\n2. 000456")
                    self.assertEqual(report["Claims"]["P2"].number_format, "@")
                    self.assertEqual(rows[1][headers.index("Claim ID")], "1. 000123\n2. 000456")
                    self.assertEqual(report["Claims"]["W2"].number_format, "@")
                    self.assertEqual(rows[1][headers.index("Billed Amount")], "1. $100.00\n2. $100.00")
                    self.assertEqual(rows[1][headers.index("Visit/Supply/Expense")], 100)
                    self.assertEqual(rows[1][headers.index("Paid Amount")], "1. $80.00\n2. ")
                    self.assertEqual(rows[1][headers.index("Check Number")], "1. 0000123\n2. 0000123")
                    self.assertEqual(report["Claims"]["AD2"].number_format, "@")
                    self.assertEqual(rows[1][headers.index("AutomationStatus")], "Pending")
                    self.assertEqual(report["Claims"]["A1"].fill.fgColor.rgb, "003D5A73")
                    self.assertEqual(report["Claims"]["A1"].font.name, "Calibri")
                    self.assertEqual(report["Claims"].column_dimensions["AA"].width, 60)
                    summary_values = dict(report["Summary"].iter_rows(min_row=3, values_only=True))
                    self.assertEqual(summary_values["Claims returned"], 2)
                    self.assertEqual(rows[2][headers.index("Sr #")], 2)
                    self.assertEqual(rows[2][headers.index("InvoiceNumber")], "INV-02")
                    self.assertEqual(rows[2][headers.index("AutomationStatus")], "Not Found")
                    self.assertIsNone(rows[2][headers.index("Claim ID")])
                    self.assertEqual(rows[2][headers.index("Claim Status")], "Not Found")
                    self.assertEqual(rows[2][headers.index("Pay. Status")], "Not Found")
                    self.assertIn("No portal claims", rows[2][headers.index("Note(s)/ Follow Up")])
                    available = {"Sr #", "Office", "Contract", "VisitDate", "Visit/Supply/Expense", "AutomationStatus", "PatientName", "InvoiceNumber", "Current Reason for Non - Payment",
                                 "Invoice Date", "Amount", "PaidAmount", "Adj.", "Balance", "Pay. Status", "TRN", "StartDate", "EndDate", "Claim ID", "Billed Amount", "Paid Amount",
                                 "AltPatientID", "Denial Reason", "Finalized Date", "Claim Status", "Check Number", "Check Date", "Note(s)/ Follow Up"}
                    for row in rows[1:]:
                        for header, value in zip(headers, row):
                            if header not in available:
                                self.assertIsNone(value, header)
                    self.assertEqual(report["Claims"].auto_filter.ref, "A1:AH4")
                    self.assertEqual(report.sheetnames, ["Claims", "Summary", "Results"])
                    self.assertEqual(report.active.title, "Claims")
                    self.assertEqual(report["Results"].sheet_state, "hidden")
                    self.assertEqual(report["Claims"].freeze_panes, "A2")
                    self.assertTrue(report["Claims"]["O2"].alignment.wrap_text)
                    for row in report["Claims"].iter_rows():
                        for cell in row:
                            for edge in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom):
                                self.assertEqual(edge.style, "thin")
                                self.assertEqual(edge.color.rgb, "00D9D9D9")
                    self.assertGreater(report["Claims"].row_dimensions[2].height, 70)
                    self.assertEqual(len(report["Claims"].conditional_formatting), 1)
                finally:
                    report.close()
            finally:
                restored.close()

    def test_not_found_saved_export_uses_input_patient_and_invoice_data(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            source = Path(temp) / "not-found-patient.xlsx"
            book = Workbook()
            book.active.append([
                "First Name", "Last Name", "Patient ID", "DOB", "Insurance ID",
                "Invoice No", "Invoice Date", "Statement Period Start Date",
                "Statement Period End Date", "Invoice Amount", "Balance Amount",
                "Invoice Notes", "Comments", "ICN", "Paid Amount", "Check Number",
            ])
            book.active.append([
                "TEST", "PATIENT", "000042", date(1970, 1, 2), "000987",
                "000123", date(2026, 6, 1), date(2026, 5, 1), date(2026, 5, 31),
                0, 0, "=Original invoice note", "Original follow-up comment",
                "SOURCE-ICN", 500, "SOURCE-CHECK",
            ])
            book.active["L2"].data_type = "s"
            book.save(source)
            book.close()
            original = source.read_bytes()
            app = create_application()
            app.automation.office = "Arizona Home Care"
            paths = ApplicationPaths("not-found-test", "not-found-test", Path(temp) / "state")

            def session():
                return WorkbookSession(app.automation.spec.workbook, app.automation.spec.variants[0],
                                       plugin_id=app.spec.id, app_paths=paths, output_policy=app.spec.output)

            current = session()
            try:
                current.prepare(str(source))
                batch, _ = current.begin_batch(10)
                record = current.claim_next(batch)
                with patch("waystar_claim_checker.application.ClaimsPage") as page:
                    page.return_value.search.return_value = []
                    result = app.automation.process_row(record, SimpleNamespace(cancellation=CancellationToken()))
                    page.return_value.search.assert_called_once_with("000123")
                self.assertEqual(result.values["claim_rows"], [])
                current.commit(record, result)
                current.finish_batch()
                run_id = current.run.id
            finally:
                current.close()

            restored = session()
            try:
                restored.run = restored.store.get_run(run_id)
                report = load_workbook(restored.export_report(), data_only=False)
                try:
                    rows = list(report["Claims"].iter_rows(values_only=True))
                    row = dict(zip(rows[0], rows[1]))
                    self.assertEqual(row["PatientName"], "TEST PATIENT")
                    self.assertEqual(row["InvoiceNumber"], "000123")
                    self.assertEqual(row["AltPatientID"], "000987")
                    self.assertEqual(row["DOB"].date(), date(1970, 1, 2))
                    self.assertEqual(row["Invoice Date"].date(), date(2026, 6, 1))
                    self.assertEqual(row["StartDate"].date(), date(2026, 5, 1))
                    self.assertEqual(row["EndDate"].date(), date(2026, 5, 31))
                    self.assertEqual(row["VisitDate"], "05/01/2026 - 05/31/2026")
                    for header in ("Visit/Supply/Expense", "Amount", "Billed Amount", "Balance"):
                        self.assertEqual(row[header], 0, header)
                    for header in ("Claim Status", "Pay. Status", "AutomationStatus"):
                        self.assertEqual(row[header], "Not Found", header)
                    self.assertIn("=Original invoice note", row["Note(s)/ Follow Up"])
                    self.assertIn("Original follow-up comment", row["Note(s)/ Follow Up"])
                    self.assertEqual(report["Claims"]["Q2"].data_type, "s")
                    for header in ("PaidAmount", "Adj.", "TRN", "Claim ID", "Paid Amount",
                                   "Finalized Date", "Check Number", "Check Date", "LastError"):
                        self.assertIsNone(row[header], header)
                    summary = dict(report["Summary"].iter_rows(min_row=3, values_only=True))
                    self.assertEqual(summary["Claims returned"], 0)
                    self.assertEqual(summary["Invoices checked"], 1)
                    original_rows = list(report["Results"].iter_rows(values_only=True))
                    original_row = dict(zip(original_rows[0], original_rows[1]))
                    self.assertEqual(original_row["Comments"], "Original follow-up comment")
                finally:
                    report.close()
                self.assertEqual(source.read_bytes(), original)
            finally:
                restored.close()

    def test_not_found_invoice_only_and_full_name_aliases(self):
        fixtures = (
            (["Invoice No"], ["000123"], None, None),
            (["Invoice Number"], ["000123"], None, None),
            (["InvoiceNumber"], ["000123"], None, None),
            (["Invoice No", "Full Name", "Patient ID"],
             ["000123", "SYNTHETIC FULL NAME", "000042"], "SYNTHETIC FULL NAME", "000042"),
        )
        for headers, values, patient_name, patient_id in fixtures:
            with self.subTest(headers=headers), tempfile.TemporaryDirectory(dir=ROOT) as temp:
                source = Path(temp) / "not-found.xlsx"
                book = Workbook()
                book.active.append(headers)
                book.active.append(values)
                book.save(source)
                book.close()
                app = create_application()
                paths = ApplicationPaths("not-found-aliases", "not-found-aliases", Path(temp) / "state")
                session = WorkbookSession(app.automation.spec.workbook, app.automation.spec.variants[0],
                                          plugin_id=app.spec.id, app_paths=paths, output_policy=app.spec.output)
                try:
                    session.prepare(str(source))
                    batch, _ = session.begin_batch(10)
                    record = session.claim_next(batch)
                    with patch("waystar_claim_checker.application.ClaimsPage") as page:
                        page.return_value.search.return_value = []
                        result = app.automation.process_row(record, SimpleNamespace(cancellation=CancellationToken()))
                    session.commit(record, result)
                    session.finish_batch()
                    report = load_workbook(session.export_report(), read_only=True, data_only=True)
                    try:
                        rows = list(report["Claims"].iter_rows(values_only=True))
                        row = dict(zip(rows[0], rows[1]))
                        self.assertEqual(row["InvoiceNumber"], "000123")
                        self.assertEqual(row["PatientName"], patient_name)
                        self.assertEqual(row["AltPatientID"], patient_id)
                        self.assertEqual(row["Claim Status"], "Not Found")
                        self.assertEqual(row["AutomationStatus"], "Not Found")
                        self.assertIsNone(row["Amount"])
                        self.assertIsNone(row["Claim ID"])
                    finally:
                        report.close()
                finally:
                    session.close()

    def test_found_claim_export_does_not_fill_portal_fields_from_source(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            source = Path(temp) / "found-with-different-input.xlsx"
            book = Workbook()
            book.active.append(["Invoice No", "Full Name", "DOB", "Invoice Date", "Invoice Amount",
                                "Statement Period Start Date", "Statement Period End Date", "Comments"])
            book.active.append(["INV-01", "INPUT PATIENT", date(1970, 1, 2), date(2025, 1, 1),
                                999, date(2025, 1, 1), date(2025, 1, 31), "Input note"])
            book.save(source)
            book.close()
            app = create_application()
            paths = ApplicationPaths("found-source-test", "found-source-test", Path(temp) / "state")
            session = WorkbookSession(app.automation.spec.workbook, app.automation.spec.variants[0],
                                      plugin_id=app.spec.id, app_paths=paths, output_policy=app.spec.output)
            try:
                session.prepare(str(source))
                batch, _ = session.begin_batch(10)
                record = session.claim_next(batch)
                with patch("waystar_claim_checker.application.ClaimsPage") as page:
                    page.return_value.search.return_value = [claim(invoice="INV-01", patient_full_name="",
                                                                  notes="", received_date=None)]
                    result = app.automation.process_row(record, SimpleNamespace(cancellation=CancellationToken()))
                session.commit(record, result)
                session.finish_batch()
                report = load_workbook(session.export_report(), read_only=True, data_only=True)
                try:
                    rows = list(report["Claims"].iter_rows(values_only=True))
                    row = dict(zip(rows[0], rows[1]))
                    self.assertIsNone(row["PatientName"])
                    self.assertIsNone(row["DOB"])
                    self.assertIsNone(row["Invoice Date"])
                    self.assertIsNone(row["Note(s)/ Follow Up"])
                    self.assertEqual(row["Amount"], 999)
                    self.assertEqual(row["Billed Amount"], 100)
                    self.assertEqual(row["VisitDate"], "05/01/2026 - 05/31/2026")
                    self.assertEqual(row["Claim Status"], "Paid")
                    self.assertEqual(row["AutomationStatus"], "Done")
                finally:
                    report.close()
            finally:
                session.close()

    def test_optional_claim_output_preserves_blank_and_zero_payments(self):
        missing = claim_output(claim(paid=None, check_number="", check_date=None,
                                     claim_id="", notes="Portal follow-up note"))
        self.assertIsNone(missing["paid_amount"])
        self.assertIsNone(missing["adjustment"])
        self.assertIsNone(missing["balance"])
        self.assertIsNone(missing["check_date"])
        self.assertIsNone(missing["finalized_date"])
        self.assertEqual(missing["claim_id"], "")
        self.assertEqual(missing["check_number"], "")
        self.assertEqual(missing["notes"], "Portal follow-up note")
        self.assertEqual(claim_output(claim(paid=Decimal("0.00")))["paid_amount"], Decimal("0.00"))
        for billed, paid, expected in (("100.00", "0.00", "100.00"),
                                      ("100.00", "100.00", "0.00"),
                                      ("100.00", "120.00", "-20.00")):
            output = claim_output(claim(billed=Decimal(billed), paid=Decimal(paid)))
            self.assertEqual(output["adjustment"], Decimal(expected))
            self.assertEqual(output["balance"], Decimal(expected))
        self.assertIsNone(claim_output(claim(billed=None))["balance"])

    def test_different_pcn_stops_without_exporting_wrong_claim(self):
        with patch("waystar_claim_checker.application.ClaimsPage") as page:
            page.return_value.search.return_value = [claim(invoice="OTHER")]
            with self.assertRaises(PortalAutomationError) as caught:
                Automation().process_row(RowRecord(0, "INV-01", input_values()),
                                         SimpleNamespace(cancellation=CancellationToken()))
        self.assertEqual(caught.exception.reason, "pcn_mismatch")

    def test_missing_pcn_skips_without_search(self):
        for pcn in (None, "", "   "):
            with self.subTest(pcn=pcn):
                with patch("waystar_claim_checker.application.ClaimsPage") as page:
                    result = Automation().process_row(
                        RowRecord(0, "", input_values(invoice_no=pcn)),
                        SimpleNamespace(cancellation=CancellationToken()))
                page.assert_not_called()
                self.assertEqual(result.disposition, RowDisposition.SKIPPED)
                self.assertEqual(result.values["claim_rows"], [])
                self.assertIsNone(result.comment)
                self.assertEqual(result.diagnostic_tags["branch"], "missing_pcn")

    def test_live_export_keeps_source_and_comments(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            source = Path(temp) / "synthetic.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Invoice No", "Insurance ID", "Statement Period Start Date",
                          "Statement Period End Date", "Invoice Amount", "Comments", "ICN"])
            v = input_values()
            sheet.append([v[k] for k in ("invoice_no", "insurance_id", "service_from", "service_to", "invoice_amount", "comments", "icn")])
            book.save(source)
            original = source.read_bytes()
            app = create_application()
            paths = ApplicationPaths("waystar-test", "waystar-test", Path(temp) / "state")
            session = WorkbookSession(app.automation.spec.workbook, app.automation.spec.variants[0],
                                      plugin_id=app.spec.id, app_paths=paths, output_policy=app.spec.output)
            try:
                session.prepare(str(source))
                batch, count = session.begin_batch(10)
                self.assertEqual(count, 1)
                record = session.claim_next(batch)
                with patch("waystar_claim_checker.application.ClaimsPage") as page:
                    page.return_value.search.return_value = [claim(invoice="INV-01")]
                    result = app.automation.process_row(record, SimpleNamespace(cancellation=CancellationToken()))
                session.commit(record, result)
                session.finish_batch()
                report = load_workbook(session.export_report(), data_only=True)
                try:
                    rows = list(report["Results"].iter_rows(values_only=True))
                    headers = rows[0]
                    self.assertEqual(rows[1][headers.index("Comments")], v["comments"])
                    self.assertIn("PCN search complete", rows[1][headers.index("Claim Checker Comments")])
                    self.assertIn("Check Date", headers)
                finally:
                    report.close()
                self.assertEqual(source.read_bytes(), original)
            finally:
                session.close()


if __name__ == "__main__":
    unittest.main()
