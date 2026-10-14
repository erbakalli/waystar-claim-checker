"""Memory-bounded Summary and Results workbook generation."""

from __future__ import annotations

from dataclasses import dataclass
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import os
from pathlib import Path
import re
import textwrap
from typing import Any, Iterator, Mapping
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.worksheet import Worksheet

from .api import ColumnFormat, ColumnSpec, WorkbookSpec
from .run_store import PersistedRow, RunInfo, RunStore
from .spreadsheet_source import (
    SourceLayout,
    find_column_index,
    iter_source_rows,
    normalize_header,
)


class ReportError(RuntimeError):
    pass


@dataclass(frozen=True)
class ReportColumn:
    header: str
    source_index: int | None
    spec: ColumnSpec | None


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
ALT_FILL = PatternFill("solid", fgColor="F3F6FA")
SUMMARY_FILL = PatternFill("solid", fgColor="D9EAF7")
PAID_FILL = PatternFill("solid", fgColor="E2F0D9")
PENDING_FILL = PatternFill("solid", fgColor="FFF2CC")
DENIED_FILL = PatternFill("solid", fgColor="FCE4D6")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TABLE_BORDER = Border(left=Side(style="thin", color="D9D9D9"),
                      right=Side(style="thin", color="D9D9D9"),
                      top=Side(style="thin", color="D9D9D9"),
                      bottom=Side(style="thin", color="D9D9D9"))
BODY_FONT = Font(name="Arial", size=10, color="1F2937")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1F2937")


def _value_cell(sheet: Any, value: Any) -> Any:
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        # Source and portal text are values, including literal leading '='.
        cell.data_type = "s"
    return cell


def _matching_spec(header: Any, columns: tuple[ColumnSpec, ...]) -> ColumnSpec | None:
    normalized = normalize_header(header)
    for column in columns:
        accepted = {normalize_header(value) for value in column.accepted_headers}
        if normalized in accepted or any(token.lower() in normalized for token in column.contains):
            return column
    return None


def _report_columns(headers: tuple[Any, ...], spec: WorkbookSpec) -> list[ReportColumn]:
    provisioned = tuple(column for column in spec.columns if column.provision and not column.anchor)
    kept: list[ReportColumn] = []
    for index, header in enumerate(headers):
        matched_output = _matching_spec(header, provisioned)
        if matched_output is not None:
            continue
        kept.append(
            ReportColumn(
                header=str(header or f"Column {index + 1}"),
                source_index=index,
                spec=_matching_spec(header, spec.columns),
            )
        )

    anchor_position = next(
        (
            index
            for index, column in enumerate(kept)
            if column.spec is not None and column.spec.key == spec.anchor.key
        ),
        None,
    )
    if anchor_position is None:
        if not spec.anchor.provision:
            raise ReportError(f"Column not found: {spec.anchor.header}")
        kept.append(ReportColumn(spec.anchor.header, None, spec.anchor))
        anchor_position = len(kept) - 1
    additions = [
        ReportColumn(column.header, find_column_index(headers, column), column)
        for column in provisioned
    ]
    return kept[: anchor_position + 1] + additions + kept[anchor_position + 1 :]


def _parse_currency(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    try:
        amount = Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return value
    if negative:
        amount = -amount
    return float(amount)


def _parse_date(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (date, datetime)):
        return value
    text = str(value).strip()
    for pattern in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return value


def _typed(value: Any, column: ReportColumn) -> Any:
    if column.spec is None:
        return value
    if column.spec.format == ColumnFormat.CURRENCY:
        return _parse_currency(value)
    if column.spec.format == ColumnFormat.DATE:
        return _parse_date(value)
    if column.spec.format in {ColumnFormat.TEXT, ColumnFormat.LONG_TEXT}:
        return None if value is None else str(value)
    return value


def _style_body_cell(cell: WriteOnlyCell, column: ReportColumn, alternate: bool) -> None:
    cell.font = BODY_FONT
    cell.alignment = Alignment(
        vertical="top" if column.spec and column.spec.format == ColumnFormat.LONG_TEXT else "center",
        horizontal=(
            "right"
            if column.spec and column.spec.format == ColumnFormat.CURRENCY
            else "left"
        ),
        wrap_text=bool(column.spec and column.spec.format == ColumnFormat.LONG_TEXT),
    )
    if alternate:
        cell.fill = ALT_FILL
    if column.spec:
        if column.spec.format == ColumnFormat.CURRENCY:
            cell.number_format = '"$"#,##0.00'
        elif column.spec.format == ColumnFormat.DATE:
            cell.number_format = "mm/dd/yyyy"
        elif column.spec.format == ColumnFormat.TEXT:
            cell.number_format = "@"


def _status_fill(value: Any) -> PatternFill | None:
    text = " ".join(str(value or "").lower().split())
    if "denied" in text or "reject" in text:
        return DENIED_FILL
    if "return to provider" in text:
        return PENDING_FILL
    if text in {"in process", "in review", "finalized", "acknowledged"}:
        return PENDING_FILL
    if any(token in text for token in ("paid", "approved", "resolved")):
        return PAID_FILL
    return None


def _detail_row_height(values: list[Any], columns: list[ReportColumn], widths: dict[str, float]) -> float:
    lines = 1
    for value, column in zip(values, columns):
        if not isinstance(value, str):
            continue
        usable = max(8, int(widths[column.spec.key]))
        wrapped = sum(max(1, len(textwrap.wrap(line, width=usable))) for line in value.splitlines())
        lines = max(lines, wrapped)
    return min(409, max(32, lines * 13 + 8))


def _group_detail_records(records, sheet_spec):
    """Keep claim numbering aligned across the columns of one invoice."""
    combined = {}
    for column in sheet_spec.columns:
        values = [record.get(column.key) for record in records]
        if not any(value is not None and value != "" for value in values):
            continue
        if (len(records) == 1 or (column.key not in sheet_spec.grouped_columns
                                  and all(value == values[0] for value in values))):
            combined[column.key] = next(value for value in values if value is not None and value != "")
            continue
        lines = []
        for number, value in enumerate(values, 1):
            if value is None or value == "":
                text = ""
            elif column.format == ColumnFormat.CURRENCY:
                text = f"${Decimal(str(value)):,.2f}"
            elif column.format == ColumnFormat.DATE:
                parsed = _parse_date(value)
                text = parsed.strftime("%m/%d/%Y") if isinstance(parsed, (date, datetime)) else str(value)
            else:
                text = str(value)
            lines.append(f"{number}. {text}")
        combined[column.key] = "\n".join(lines)
    return combined


def _next_persisted(
    iterator: Iterator[PersistedRow], current: PersistedRow | None, source_row: int
) -> tuple[PersistedRow | None, PersistedRow | None]:
    while current is not None and current.source_row < source_row:
        current = next(iterator, None)
    return (current if current is not None and current.source_row == source_row else None), current


def generate_report(
    *,
    store: RunStore,
    run: RunInfo,
    spec: WorkbookSpec,
    statistics: dict[str, tuple[str, int]],
) -> Path:
    output = Path(run.output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    if not os.access(output.parent, os.W_OK):
        raise ReportError(f"The report folder is not writable: {output.parent}")
    temp = output.with_name(f".{output.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    headers = tuple(store.headers(run.id))
    columns = _report_columns(headers, spec)
    layout = SourceLayout(
        extension=run.source_ext,
        sheet_name=run.sheet_name,
        header_row=run.header_row,
        first_col=run.first_col,
        headers=headers,
    )
    if run.source_ext == ".csv":
        from .spreadsheet_source import _csv_settings  # localized internal reuse

        encoding, delimiter = _csv_settings(Path(run.snapshot_path))
        layout = SourceLayout(
            extension=run.source_ext,
            sheet_name=run.sheet_name,
            header_row=run.header_row,
            first_col=run.first_col,
            headers=headers,
            encoding=encoding,
            delimiter=delimiter,
        )

    counts = store.status_counts(run.id)
    current_run = store.get_run(run.id)
    workbook = Workbook(write_only=True)
    primary_sheets = {s.name: workbook.create_sheet(s.name) for s in spec.result_sheets if s.primary}
    if spec.hide_source_results:
        workbook.active = 0
    summary = workbook.create_sheet("Summary")
    summary.sheet_view.showGridLines = False
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 44
    title = _value_cell(summary, f"{run.automation_name} Results")
    title.font = TITLE_FONT
    summary.append([title])
    summary.append([])
    completed = counts.get("RESOLVED", 0) + counts.get("SKIPPED", 0)
    summary_rows = [
        ("Automation", run.automation_name),
        ("Source file", Path(run.source_path).name),
        ("Worksheet", run.sheet_name),
        ("Run status", current_run.status.title()),
        ("Created", run.created_at),
        ("Last updated", current_run.updated_at),
        ("Total rows", run.row_count),
        ("Completed", completed),
        ("Resolved", counts.get("RESOLVED", 0)),
        ("Skipped", counts.get("SKIPPED", 0)),
        ("Retry later", counts.get("RETRY_LATER", 0)),
        ("Pending", counts.get("PENDING", 0) + counts.get("IN_PROGRESS", 0)),
    ]
    if current_run.completed_at:
        summary_rows.insert(6, ("Completed", current_run.completed_at))
    summary_rows.extend((label, value) for _key, (label, value) in statistics.items())
    if spec.compact_summary:
        claim_count = 0
        statuses = Counter()
        for persisted in store.iter_rows(run.id):
            for sheet in spec.result_sheets:
                for record in persisted.result.get(sheet.result_key, []):
                    claim_count += 1
                    status = record.get("claim_status", sheet.missing_values.get("claim_status"))
                    if status:
                        statuses[str(status)] += 1
        summary.column_dimensions["A"].width = 34
        summary.column_dimensions["B"].width = 44
        summary_rows = [
            ("Source file", Path(run.source_path).name),
            ("Invoices in workbook", run.row_count),
            ("Invoices checked", completed),
            ("Invoices remaining", run.row_count - completed),
            ("Claims returned", claim_count),
            ("", ""),
            ("Claim Status", "Claims"),
            *sorted(statuses.items()),
            ("", ""),
            ("Status colors", "Rejected / denied: red. Return to Provider / pending: amber. Paid / approved: green."),
        ]
        if "Not captured" in statuses:
            summary_rows.append(("Earlier records", "Claim Status was not saved in earlier runs. New claim checks include the portal status."))
    for summary_index, (label, value) in enumerate(summary_rows, start=3):
        left = _value_cell(summary, label)
        left.font = Font(name="Arial", size=10, bold=True, color="1F2937")
        left.fill = SUMMARY_FILL
        left.alignment = Alignment(vertical="center")
        right = _value_cell(summary, value)
        right.font = BODY_FONT
        right.alignment = Alignment(vertical="top" if spec.compact_summary else "center",
                                    wrap_text=spec.compact_summary)
        if spec.compact_summary:
            lines = max(1, len(textwrap.wrap(str(value), width=35)))
            summary.row_dimensions[summary_index].height = max(26, lines * 15 + 10)
        summary.append([left, right])

    results = workbook.create_sheet("Results")
    if spec.hide_source_results:
        results.sheet_state = "hidden"
    results.sheet_view.showGridLines = False
    results.freeze_panes = "A2"
    for index, column in enumerate(columns, start=1):
        if column.spec and column.spec.format == ColumnFormat.LONG_TEXT:
            width = 48
        elif column.spec and column.spec.format == ColumnFormat.CURRENCY:
            width = 16
        elif column.spec and column.spec.format == ColumnFormat.DATE:
            width = 14
        else:
            width = min(32, max(12, len(column.header) + 3))
        results.column_dimensions[get_column_letter(index)].width = width
    header_cells = []
    for column in columns:
        cell = _value_cell(results, column.header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cells.append(cell)
    results.append(header_cells)

    persisted_iterator = iter(store.iter_rows(run.id))
    persisted_current = next(persisted_iterator, None)
    written_rows = 0
    status_index = next(
        (
            index
            for index, column in enumerate(columns)
            if column.spec is not None and "status" in column.spec.key
        ),
        None,
    )
    for source_row, source_values in iter_source_rows(run.snapshot_path, layout):
        if written_rows >= run.row_count:
            break
        persisted, persisted_current = _next_persisted(
            persisted_iterator, persisted_current, source_row
        )
        overlay = dict(persisted.result) if persisted is not None else {}
        comment = persisted.comment if persisted is not None else None
        row_values: list[Any] = []
        for column in columns:
            original = (
                source_values[column.source_index]
                if column.source_index is not None and column.source_index < len(source_values)
                else None
            )
            value = original
            if column.spec is not None:
                if column.spec.key == spec.anchor.key and comment is not None:
                    value = comment
                elif column.spec.key in overlay:
                    value = overlay[column.spec.key]
            row_values.append(_typed(value, column))

        cells = []
        alternate = written_rows % 2 == 1
        for index, (value, column) in enumerate(zip(row_values, columns)):
            cell = _value_cell(results, value)
            _style_body_cell(cell, column, alternate)
            if status_index == index:
                fill = _status_fill(value)
                if fill is not None:
                    cell.fill = fill
            cells.append(cell)
        results.append(cells)
        written_rows += 1
    results.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{written_rows + 1}"

    detail_counts = {}
    for sheet_spec in spec.result_sheets:
        if ((sheet_spec.name in workbook.sheetnames and sheet_spec.name not in primary_sheets)
                or not sheet_spec.columns):
            raise ReportError("Result sheets require unique names and at least one column")
        detail = primary_sheets.get(sheet_spec.name)
        if detail is None:
            detail = workbook.create_sheet(sheet_spec.name)
        detail.freeze_panes = sheet_spec.freeze_panes
        detail.sheet_view.showGridLines = False
        detail.sheet_view.zoomScale = 85
        detail.sheet_properties.pageSetUpPr.fitToPage = True
        detail.page_setup.orientation = "landscape"
        detail.page_setup.paperSize = Worksheet.PAPERSIZE_A3
        detail.page_setup.fitToWidth = 1
        detail.page_setup.fitToHeight = 0
        detail.print_title_rows = "1:1"
        detail.row_dimensions[1].height = 32
        detail_columns = [ReportColumn(c.header, None, c) for c in sheet_spec.columns]
        widths = {c.key: sheet_spec.column_widths.get(c.key,
                  72 if c.format == ColumnFormat.LONG_TEXT else
                  14 if c.format == ColumnFormat.DATE else
                  16 if c.format == ColumnFormat.CURRENCY else 24) for c in sheet_spec.columns}
        detail_headers = []
        for index, column in enumerate(detail_columns, 1):
            detail.column_dimensions[get_column_letter(index)].width = widths[column.spec.key]
            cell = _value_cell(detail, column.header)
            cell.font = Font(name=sheet_spec.font_name, size=sheet_spec.font_size, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=sheet_spec.header_color)
            if sheet_spec.cell_borders:
                cell.border = TABLE_BORDER
            cell.alignment = Alignment(horizontal="center" if sheet_spec.grouped_columns else "left",
                                       vertical="center", wrap_text=True)
            detail_headers.append(cell)
        detail.append(detail_headers)
        shared_values = {}
        if sheet_spec.shared_columns:
            for persisted in store.iter_rows(run.id):
                for key in sheet_spec.shared_columns:
                    value = persisted.result.get(key)
                    if value is not None and value != "":
                        if key in shared_values and shared_values[key] != value:
                            raise ReportError(f"Result sheet has conflicting values for shared column {key}")
                        shared_values[key] = value
        count = 0
        source_iterator = iter(iter_source_rows(run.snapshot_path, layout)) if sheet_spec.empty_result_factory else None
        source_current = next(source_iterator, None) if source_iterator is not None else None
        for persisted in store.iter_rows(run.id):
            records = persisted.result.get(sheet_spec.result_key, [])
            if not isinstance(records, list):
                raise ReportError("Result sheet records must be a list")
            if any(not isinstance(record, dict) for record in records):
                raise ReportError("Result sheet rows must be mappings")
            if sheet_spec.grouped_columns:
                combined = _group_detail_records(records, sheet_spec)
                combined["automation_status"] = {
                    "RESOLVED": "Done", "SKIPPED": "Skipped", "RETRY_LATER": "Retry later",
                }.get(persisted.status, persisted.status.title())
                if (persisted.status == "RESOLVED" and any(
                        str(record.get("claim_status", "")).strip().casefold()
                        in {status.casefold() for status in sheet_spec.pending_claim_statuses}
                        for record in records)):
                    combined["automation_status"] = "Pending"
                if not records and persisted.status == "RESOLVED" and sheet_spec.empty_result_factory:
                    while source_current is not None and source_current[0] < persisted.source_row:
                        source_current = next(source_iterator, None)
                    if source_current is None or source_current[0] != persisted.source_row:
                        raise ReportError("Saved input row is missing from the source snapshot")
                    source_data = dict(zip(headers, source_current[1]))
                    fallback = sheet_spec.empty_result_factory(source_data, persisted.result)
                    if not isinstance(fallback, Mapping):
                        raise ReportError("Empty result factory must return a mapping")
                    combined.update(fallback)
                combined.setdefault("pcn", persisted.result.get("pcn"))
                if persisted.status == "RETRY_LATER":
                    combined["last_error"] = persisted.comment or persisted.result.get("checker_comments")
                records = [combined]
            for record in records:
                if not isinstance(record, dict):
                    raise ReportError("Result sheet rows must be mappings")
                cells = []
                row_values = []
                for column in detail_columns:
                    value = (shared_values.get(column.spec.key) if column.spec.key in sheet_spec.shared_columns
                             else record.get(column.spec.key, sheet_spec.missing_values.get(column.spec.key)))
                    if (column.spec.key in sheet_spec.processed_only_columns
                            and persisted.status not in {"RESOLVED", "SKIPPED", "RETRY_LATER"}):
                        value = None
                    if column.spec.key == sheet_spec.serial_column:
                        value = count + 1
                    row_values.append(value if isinstance(value, str) and "\n" in value else _typed(value, column))
                detail.row_dimensions[count + 2].height = max(
                    60 if sheet_spec.grouped_columns else 32,
                    _detail_row_height(row_values, detail_columns, widths))
                for value, column in zip(row_values, detail_columns):
                    cell = _value_cell(detail, value)
                    _style_body_cell(cell, column, count % 2 == 1 and not sheet_spec.grouped_columns)
                    cell.font = Font(name=sheet_spec.font_name, size=sheet_spec.font_size, color="1F2937")
                    if sheet_spec.cell_borders:
                        cell.border = TABLE_BORDER
                    if isinstance(value, str) and "\n" in value:
                        cell.number_format = "@"
                    cell.alignment = Alignment(vertical="top", wrap_text=True,
                                               horizontal="right" if column.spec.format == ColumnFormat.CURRENCY else "left")
                    cells.append(cell)
                detail.append(cells)
                count += 1
        if source_iterator is not None:
            source_iterator.close()
        detail.auto_filter.ref = f"A1:{get_column_letter(len(detail_columns))}{count + 1}"
        status_column = next((i for i, c in enumerate(sheet_spec.columns, 1) if c.key == "claim_status"), None)
        if status_column and count:
            letter = get_column_letter(status_column)
            for keywords, fill in ((("reject", "denied"), DENIED_FILL),
                                   (("return to provider", "pending"), PENDING_FILL),
                                   (("paid", "approved"), PAID_FILL)):
                tests = ",".join(f'ISNUMBER(SEARCH("{word}",{letter}2))' for word in keywords)
                detail.conditional_formatting.add(f"{letter}2:{letter}{count + 1}",
                    FormulaRule(formula=[f"OR({tests})"], fill=fill, stopIfTrue=True))
        detail_counts[sheet_spec.name] = count

    try:
        workbook.save(temp)
        check = load_workbook(temp, read_only=True, data_only=True)
        try:
            expected_sheets = [*primary_sheets, "Summary", "Results",
                               *(name for name in detail_counts if name not in primary_sheets)]
            if check.sheetnames != expected_sheets:
                raise ReportError("The generated report has an unexpected worksheet layout")
            result_sheet = check["Results"]
            checked_rows = sum(1 for _row in result_sheet.iter_rows(values_only=True))
            if checked_rows != written_rows + 1:
                raise ReportError("The generated report row count could not be verified")
            checked_headers = [cell for cell in next(result_sheet.iter_rows(values_only=True))]
            if checked_headers != [column.header for column in columns]:
                raise ReportError("The generated report headers could not be verified")
            for sheet_spec in spec.result_sheets:
                detail_rows = check[sheet_spec.name].iter_rows(values_only=True)
                if list(next(detail_rows)) != [c.header for c in sheet_spec.columns]:
                    raise ReportError("Result sheet headers could not be verified")
                if sum(1 for _ in detail_rows) != detail_counts[sheet_spec.name]:
                    raise ReportError("Result sheet row count could not be verified")
        finally:
            check.close()
        os.replace(temp, output)
    except PermissionError as exc:
        if spec.save_copy_if_locked and temp.is_file():
            alternate = output.with_name(f"{output.stem} - Updated {uuid.uuid4().hex[:8]}.xlsx")
            try:
                os.replace(temp, alternate)
                output = alternate
            except OSError as copy_error:
                raise ReportError(f"Could not save an updated copy of {output.name}: {copy_error}") from copy_error
        else:
            raise ReportError(
                f"Could not replace {output.name}. Close the report if it is open, then use Export Results."
            ) from exc
    except OSError as exc:
        raise ReportError(f"Could not save report: {exc}") from exc
    finally:
        try:
            if temp.exists():
                temp.unlink()
        except OSError:
            pass
    return output
