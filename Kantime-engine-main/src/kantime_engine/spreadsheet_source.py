"""Streaming XLSX/CSV inspection and row iteration."""

from __future__ import annotations

import csv
import codecs
from dataclasses import dataclass
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import os
from typing import Any, Iterator

from openpyxl import load_workbook

from .api import ColumnSpec, VariantSpec, WorkbookSpec


class SpreadsheetSourceError(RuntimeError):
    pass


SUPPORTED_EXTENSIONS = {".xlsx", ".csv"}


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().lower().split()).rstrip(".")


def normalize_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none"} else text


def validate_extension(path: str | Path) -> str:
    extension = Path(path).suffix.lower()
    if extension in {".xls", ".xlsm"}:
        raise SpreadsheetSourceError(
            f"{extension} files are not supported. Convert the file to .xlsx or .csv first."
        )
    if extension not in SUPPORTED_EXTENSIONS:
        raise SpreadsheetSourceError("Select an .xlsx or .csv collection file.")
    return extension


@dataclass(frozen=True)
class SourceLayout:
    extension: str
    sheet_name: str
    header_row: int
    first_col: int
    headers: tuple[Any, ...]
    encoding: str | None = None
    delimiter: str | None = None


def _select_sheet(names: list[str], variant: VariantSpec) -> str:
    if variant.sheet.exact_name:
        target = normalize_header(variant.sheet.exact_name)
        for name in names:
            if normalize_header(name) == target:
                return name
        raise SpreadsheetSourceError(f"Worksheet not found: {variant.sheet.exact_name}")
    if variant.sheet.first_sheet and names:
        return names[0]
    raise SpreadsheetSourceError("No valid worksheet found")


def _trim_header(values: tuple[Any, ...] | list[Any]) -> tuple[int, tuple[Any, ...]]:
    indexes = [index for index, value in enumerate(values) if normalize_header(value)]
    if not indexes:
        raise SpreadsheetSourceError("The collection file has no header row")
    first = indexes[0]
    last = indexes[-1]
    headers = []
    for index, value in enumerate(values[first : last + 1], start=first + 1):
        headers.append(value if normalize_header(value) else f"Column {index}")
    return first + 1, tuple(headers)


def _csv_settings(path: Path) -> tuple[str, str]:
    with path.open("rb") as sample_handle:
        raw = sample_handle.read(65536)
    encoding = "utf-8-sig"
    try:
        decoder = codecs.getincrementaldecoder(encoding)()
        with path.open("rb") as handle:
            while chunk := handle.read(1024 * 1024):
                decoder.decode(chunk)
            decoder.decode(b"", final=True)
    except UnicodeDecodeError:
        encoding = "cp1252"
    sample = raw.decode(encoding)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    return encoding, delimiter


def inspect_source(path: str | Path, variant: VariantSpec) -> SourceLayout:
    resolved = Path(path).expanduser().resolve()
    extension = validate_extension(resolved)
    if extension == ".xlsx":
        try:
            workbook = load_workbook(resolved, read_only=True, data_only=True)
        except Exception as exc:
            raise SpreadsheetSourceError(f"Could not read XLSX file: {exc}") from exc
        try:
            sheet_name = _select_sheet(list(workbook.sheetnames), variant)
            worksheet = workbook[sheet_name]
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if any(normalize_header(value) for value in row):
                    first_col, headers = _trim_header(list(row))
                    return SourceLayout(
                        extension=extension,
                        sheet_name=sheet_name,
                        header_row=row_number,
                        first_col=first_col,
                        headers=headers,
                    )
            raise SpreadsheetSourceError("The selected worksheet is empty")
        finally:
            workbook.close()

    encoding, delimiter = _csv_settings(resolved)
    try:
        with resolved.open("r", encoding=encoding, newline="") as handle:
            for row_number, row in enumerate(csv.reader(handle, delimiter=delimiter), start=1):
                if any(normalize_header(value) for value in row):
                    first_col, headers = _trim_header(row)
                    return SourceLayout(
                        extension=extension,
                        sheet_name=variant.sheet.exact_name or "CSV",
                        header_row=row_number,
                        first_col=first_col,
                        headers=headers,
                        encoding=encoding,
                        delimiter=delimiter,
                    )
    except OSError as exc:
        raise SpreadsheetSourceError(f"Could not read CSV file: {exc}") from exc
    raise SpreadsheetSourceError("The CSV file is empty")


def row_has_value(values: tuple[Any, ...] | list[Any]) -> bool:
    return any(value is not None and str(value).strip() != "" for value in values)


def iter_source_rows(
    path: str | Path, layout: SourceLayout
) -> Iterator[tuple[int, tuple[Any, ...]]]:
    resolved = Path(path).expanduser().resolve()
    width = len(layout.headers)
    if layout.extension == ".xlsx":
        workbook = load_workbook(resolved, read_only=True, data_only=True)
        try:
            worksheet = workbook[layout.sheet_name]

            def generate() -> Iterator[tuple[int, tuple[Any, ...]]]:
                for row_number, row in enumerate(
                    worksheet.iter_rows(
                        min_row=layout.header_row + 1,
                        min_col=layout.first_col,
                        max_col=layout.first_col + width - 1,
                        values_only=True,
                    ),
                    start=layout.header_row + 1,
                ):
                    yield row_number, tuple(row)

            yield from generate()
        finally:
            workbook.close()
        return

    encoding = layout.encoding or "utf-8-sig"
    delimiter = layout.delimiter or ","
    with resolved.open("r", encoding=encoding, newline="") as handle:

        def generate_csv() -> Iterator[tuple[int, tuple[Any, ...]]]:
            for row_number, row in enumerate(csv.reader(handle, delimiter=delimiter), start=1):
                if row_number <= layout.header_row:
                    continue
                selected = list(row[layout.first_col - 1 : layout.first_col - 1 + width])
                selected.extend([None] * (width - len(selected)))
                yield row_number, tuple(selected)

        yield from generate_csv()


def find_column_index(headers: tuple[Any, ...], spec: ColumnSpec) -> int | None:
    accepted = {normalize_header(header) for header in spec.accepted_headers}
    for index, header in enumerate(headers):
        normalized = normalize_header(header)
        if normalized in accepted or any(token.lower() in normalized for token in spec.contains):
            return index
    return None


def map_columns(headers: tuple[Any, ...], spec: WorkbookSpec) -> dict[str, int | None]:
    return {column.key: find_column_index(headers, column) for column in spec.columns}


def schema_hash(spec: WorkbookSpec, headers: tuple[Any, ...]) -> str:
    payload = {
        "key": spec.key_column,
        "headers": [normalize_header(value) for value in headers],
        "columns": [
            {
                "key": column.key,
                "header": normalize_header(column.header),
                "aliases": [normalize_header(value) for value in column.aliases],
                "contains": list(column.contains),
                "required": column.required,
                "provision": column.provision,
                "anchor": column.anchor,
                "format": column.format.value,
            }
            for column in spec.columns
        ],
    }
    if spec.processing_revision:
        payload["processing_revision"] = spec.processing_revision
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def file_hash(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def copy_with_hash(source: str | Path, destination: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(source).open("rb") as reader, Path(destination).open("wb") as writer:
        while chunk := reader.read(1024 * 1024):
            writer.write(chunk)
            digest.update(chunk)
        writer.flush()
        os.fsync(writer.fileno())
    return digest.hexdigest()
