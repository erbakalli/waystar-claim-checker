from __future__ import annotations

import os
from pathlib import Path
import tempfile
import tracemalloc
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tests import _bootstrap  # noqa: F401

from tests.fakes import SPEC
from kantime_engine.paths import ApplicationPaths
from kantime_engine.workbook import WorkbookSession


class LargeStreamingTests(unittest.TestCase):
    def test_ten_thousand_row_xlsx_ingest_and_export_stay_memory_bounded(self):
        with tempfile.TemporaryDirectory() as temp, patch.dict(
            os.environ, {"LOCALAPPDATA": temp}
        ):
            source = Path(temp) / "large.xlsx"
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("Data")
            sheet.append(["Record ID", "Comments", "Description", "Portal Value"])
            for index in range(10_000):
                sheet.append([f"{index:08d}", None, f"Record {index}", None])
            workbook.save(source)

            session = WorkbookSession(
                SPEC.workbook,
                SPEC.variants[0],
                plugin_id=SPEC.id,
                automation_name=SPEC.name,
                app_paths=ApplicationPaths("test", "Test App", Path(temp)),
            )
            try:
                tracemalloc.start()
                session.prepare(str(source))
                _current, ingest_peak = tracemalloc.get_traced_memory()
                tracemalloc.reset_peak()
                report = session.export_report()
                _current, export_peak = tracemalloc.get_traced_memory()
                tracemalloc.stop()
                self.assertEqual(session.row_count, 10_000)
                self.assertLess(ingest_peak, 80 * 1024 * 1024)
                self.assertLess(export_peak, 96 * 1024 * 1024)
            finally:
                if tracemalloc.is_tracing():
                    tracemalloc.stop()
                session.close()

            checked = load_workbook(report, read_only=True, data_only=True)
            try:
                row_count = sum(
                    1 for _row in checked["Results"].iter_rows(values_only=True)
                )
                self.assertEqual(row_count, 10_001)
            finally:
                checked.close()


if __name__ == "__main__":
    unittest.main()
