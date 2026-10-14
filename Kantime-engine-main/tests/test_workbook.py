from __future__ import annotations

import hashlib
from dataclasses import replace
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tests import _bootstrap  # noqa: F401

from tests.fakes import FORMATTED_SPEC, SPEC
from kantime_engine.api import RowDisposition, RowResult
from kantime_engine.paths import ApplicationPaths
from kantime_engine.workbook import WorkbookError, WorkbookSession


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class WorkbookSessionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {"LOCALAPPDATA": self.temp.name})
        self.env.start()
        self.root = Path(self.temp.name)
        self.paths = ApplicationPaths("test", "Test App", self.root)
        self.variant = SPEC.variants[0]

    def tearDown(self):
        self.env.stop()
        self.temp.cleanup()

    def _xlsx(self, name: str = "collection.xlsx") -> Path:
        path = self.root / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Record ID", "Portal Value", "Other", "Comments"])
        sheet.append(["001001", "original result", 12.5, None])
        sheet.append(["001002", "pending result", 7, "keep this"])
        workbook.create_sheet("Unrelated").append(["This sheet is not exported"])
        workbook.save(path)
        workbook.close()
        return path

    def _session(self) -> WorkbookSession:
        return WorkbookSession(
            SPEC.workbook,
            self.variant,
            plugin_id=SPEC.id,
            automation_name=SPEC.name,
            app_paths=self.paths,
        )

    def test_prepare_commit_and_report_preserve_source_and_pending_values(self):
        source = self._xlsx()
        original_hash = _sha256(source)
        session = self._session()
        try:
            prepared = session.prepare(str(source))
            self.assertEqual(prepared.sheet_name, "Data")
            self.assertEqual(session.row_count, 2)
            self.assertEqual(session.statistics()["total"], ("Total Rows", 2))

            batch_id, total = session.begin_batch(1)
            self.assertEqual(total, 1)
            record = session.claim_next(batch_id)
            self.assertIsNotNone(record)
            assert record is not None
            self.assertEqual(record.index, 0)
            self.assertEqual(record.key, "001001")
            session.commit(
                record,
                RowResult(
                    RowDisposition.RESOLVED,
                    comment="done",
                    values={"portal_value": "Example Domain"},
                ),
            )
            self.assertEqual(session.finish_batch(), 1)
            report = session.export_report()
        finally:
            session.close()

        self.assertEqual(_sha256(source), original_hash)
        self.assertEqual(
            report,
            source.with_name("collection - Output") / "collection - default - Results.xlsx",
        )
        result = load_workbook(report, data_only=True)
        try:
            self.assertEqual(result.sheetnames, ["Summary", "Results"])
            values = list(result["Results"].iter_rows(values_only=True))
            self.assertEqual(
                values[0],
                ("Record ID", "Other", "Comments", "Portal Value"),
            )
            self.assertEqual(values[1], ("001001", 12.5, "done", "Example Domain"))
            self.assertEqual(
                values[2], ("001002", 7, "keep this", "pending result")
            )
            summary = dict(
                row[:2]
                for row in result["Summary"].iter_rows(values_only=True)
                if row[0]
            )
            self.assertEqual(summary["Run status"], "Paused")
            self.assertEqual(summary["Resolved"], 1)
            self.assertEqual(summary["Pending"], 1)
        finally:
            result.close()

    def test_missing_required_column_fails_before_creating_a_run(self):
        source = self.root / "missing.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Comments"])
        workbook.save(source)
        workbook.close()

        session = self._session()
        try:
            with self.assertRaisesRegex(WorkbookError, "Record ID"):
                session.prepare(str(source))
        finally:
            session.close()

    def test_export_preserves_literal_equals_text_in_headers_and_values(self):
        source = self.root / "literal-text.csv"
        source.write_text(
            "Record ID,Comments,Portal Value,=Source Header\n"
            "=SOURCE-ID,=source comment,=old value,=source note\n",
            encoding="utf-8",
        )
        original_hash = _sha256(source)
        session = self._session()
        try:
            session.prepare(str(source))
            batch_id, _ = session.begin_batch(1)
            record = session.claim_next(batch_id)
            self.assertIsNotNone(record)
            session.commit(record, RowResult(
                RowDisposition.RESOLVED, comment="=literal comment",
                values={"portal_value": "=literal portal value"},
            ))
            session.finish_batch()
            output = session.export_report()
            for data_only in (False, True):
                with self.subTest(data_only=data_only):
                    workbook = load_workbook(output, data_only=data_only)
                    try:
                        results = workbook["Results"]
                        cells = {cell.value: cell for row in results for cell in row}
                        for text in ("=Source Header", "=SOURCE-ID", "=source note",
                                     "=literal comment", "=literal portal value"):
                            self.assertIn(text, cells)
                            self.assertEqual(cells[text].data_type, "s")
                    finally:
                        workbook.close()
            self.assertEqual(_sha256(source), original_hash)
        finally:
            session.close()

    def test_csv_cp1252_delimiter_blank_rows_and_leading_zero_keys(self):
        source = self.root / "collection.csv"
        source.write_bytes(
            (
                "Record ID;Comments;Note;Portal Value\r\n"
                '0001;;"Café, East";\r\n'
                ";;;\r\n"
                "0002;;;\r\n"
                ";;;\r\n"
                ";;;\r\n"
            ).encode("cp1252")
        )
        session = self._session()
        try:
            session.prepare(str(source))
            self.assertEqual(session.row_count, 3)
            batch_id, total = session.begin_batch(3)
            self.assertEqual(total, 3)
            first = session.claim_next(batch_id)
            assert first is not None
            self.assertEqual(first.key, "0001")
            self.assertEqual(first.values["record_id"], "0001")
            session.commit(first, RowResult(RowDisposition.RESOLVED))
            blank = session.claim_next(batch_id)
            assert blank is not None
            self.assertEqual(blank.key, "")
            session.commit(blank, RowResult(RowDisposition.SKIPPED))
            last = session.claim_next(batch_id)
            assert last is not None
            self.assertEqual(last.key, "0002")
        finally:
            session.close()

    def test_csv_utf8_bom_and_quoted_fields(self):
        source = self.root / "utf8.csv"
        source.write_bytes(
            b"\xef\xbb\xbfRecord ID,Comments,Portal Value\r\n"
            b'0007,"quoted, comment",\r\n'
        )
        session = self._session()
        try:
            session.prepare(str(source))
            batch_id, _total = session.begin_batch(1)
            record = session.claim_next(batch_id)
            assert record is not None
            self.assertEqual(record.key, "0007")
            self.assertEqual(record.values["comments"], "quoted, comment")
        finally:
            session.close()

    def test_exact_fingerprint_resume_and_start_new(self):
        source = self._xlsx()
        first = self._session()
        first.prepare(str(source))
        original_run = first.run
        assert original_run is not None
        batch_id, _total = first.begin_batch(1)
        record = first.claim_next(batch_id)
        assert record is not None
        first.commit(record, RowResult(RowDisposition.RESOLVED, comment="saved"))
        first.pause()
        first.close()

        questions: list[tuple[int, int]] = []
        resumed = self._session()
        prepared = resumed.prepare(
            str(source), lambda completed, total: questions.append((completed, total)) or True
        )
        self.assertTrue(prepared.resumed)
        self.assertEqual(questions, [(1, 2)])
        self.assertEqual(resumed.run.id, original_run.id)
        resumed.pause()
        resumed.close()

        replacement = self._session()
        prepared = replacement.prepare(str(source), lambda _completed, _total: False)
        try:
            self.assertFalse(prepared.resumed)
            self.assertNotEqual(replacement.run.id, original_run.id)
            self.assertEqual(replacement.store.get_run(original_run.id).status, "ABANDONED")
        finally:
            replacement.close()

    def test_completed_or_changed_source_creates_an_isolated_run(self):
        source = self._xlsx()
        completed = self._session()
        completed.prepare(str(source))
        completed_id = completed.run.id
        batch_id, _total = completed.begin_batch(10)
        while record := completed.claim_next(batch_id):
            completed.commit(record, RowResult(RowDisposition.RESOLVED))
        self.assertEqual(completed.finish_batch(), 0)
        completed.close()

        after_complete = self._session()
        after_complete.prepare(
            str(source),
            lambda _done, _total: self.fail("completed runs must not prompt to resume"),
        )
        try:
            self.assertNotEqual(after_complete.run.id, completed_id)
            changed_from = after_complete.run.id
        finally:
            after_complete.close()

        workbook = load_workbook(source)
        workbook["Data"].append(["001003", None, 2, None])
        workbook.save(source)
        workbook.close()

        changed = self._session()
        changed.prepare(
            str(source),
            lambda _done, _total: self.fail("changed files must not prompt to resume"),
        )
        try:
            self.assertNotEqual(changed.run.id, changed_from)
            self.assertEqual(changed.row_count, 3)
        finally:
            changed.close()

    def test_interrupted_preparation_is_never_offered_for_resume(self):
        source = self._xlsx()
        interrupted = self._session()
        with patch.object(interrupted, "_ingest", side_effect=RuntimeError("crash")):
            with self.assertRaisesRegex(RuntimeError, "crash"):
                interrupted.prepare(str(source))
        row = interrupted.store.connection.execute(
            "SELECT status, is_prepared FROM runs"
        ).fetchone()
        self.assertEqual(tuple(row), ("ABANDONED", 0))
        interrupted.close()

        replacement = self._session()
        replacement.prepare(
            str(source),
            lambda _done, _total: self.fail("partial preparation must not be resumed"),
        )
        try:
            self.assertEqual(replacement.row_count, 2)
        finally:
            replacement.close()

    def test_rejects_legacy_and_macro_workbooks_with_conversion_message(self):
        for extension in (".xls", ".xlsm"):
            source = self.root / f"legacy{extension}"
            source.touch()
            session = self._session()
            try:
                with self.assertRaisesRegex(WorkbookError, r"Convert.*\.xlsx or \.csv"):
                    session.prepare(str(source))
            finally:
                session.close()

    def test_report_writes_declared_currency_date_and_text_cell_types(self):
        source = self.root / "formatted.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Invoice Number", "Comments"])
        sheet.append(["000007", None])
        workbook.save(source)
        workbook.close()

        session = WorkbookSession(
            FORMATTED_SPEC.workbook,
            FORMATTED_SPEC.variants[0],
            plugin_id=FORMATTED_SPEC.id,
            automation_name=FORMATTED_SPEC.name,
            app_paths=self.paths,
        )
        try:
            session.prepare(str(source))
            batch_id, _total = session.begin_batch(1)
            record = session.claim_next(batch_id)
            assert record is not None
            session.commit(
                record,
                RowResult(
                    RowDisposition.RESOLVED,
                    values={
                        "claim_id": "000099",
                        "billed_amount": "$1,234.56",
                        "paid_amount": "987.65",
                        "claim_status": "Paid",
                        "check_number": "000123",
                        "check_date": "09/27/2026",
                    },
                ),
            )
            session.finish_batch()
            report = session.export_report()
        finally:
            session.close()

        result = load_workbook(report, data_only=True)
        try:
            sheet = result["Results"]
            headers = [cell.value for cell in sheet[1]]
            billed = sheet.cell(2, headers.index("Billed Amount") + 1)
            check_date = sheet.cell(2, headers.index("Check Date") + 1)
            invoice = sheet.cell(2, headers.index("Invoice Number") + 1)
            claim_id = sheet.cell(2, headers.index("Claim ID") + 1)
            self.assertEqual(billed.value, 1234.56)
            self.assertEqual(billed.number_format, '"$"#,##0.00')
            self.assertEqual(check_date.value.date().isoformat(), "2026-09-27")
            self.assertEqual(check_date.number_format, "mm/dd/yyyy")
            self.assertEqual(invoice.value, "000007")
            self.assertEqual(invoice.number_format, "@")
            self.assertEqual(claim_id.value, "000099")
            self.assertEqual(claim_id.number_format, "@")
        finally:
            result.close()

    def test_locked_report_saves_updated_copy_when_enabled(self):
        source = self._xlsx("locked-copy.xlsx")
        session = self._session()
        session.spec = replace(session.spec, save_copy_if_locked=True)
        try:
            session.prepare(str(source))
            report = session.export_report()
            previous = report.read_bytes()
            batch, _ = session.begin_batch(1)
            record = session.claim_next(batch)
            session.commit(record, RowResult(RowDisposition.RESOLVED, values={"portal_value": "new result"}))
            session.finish_batch()
            real_replace = os.replace
            def locked_original(src, dst):
                if Path(dst) == report:
                    raise PermissionError("Workbook is open in Excel")
                return real_replace(src, dst)
            with patch("kantime_engine.reporting.os.replace", side_effect=locked_original):
                updated = session.export_report()
            self.assertNotEqual(updated, report)
            self.assertIn("Updated", updated.name)
            self.assertEqual(report.read_bytes(), previous)
            self.assertEqual(session.store.completed_count(session.run.id), 1)
            result = load_workbook(updated, read_only=True, data_only=True)
            try:
                rows = list(result["Results"].iter_rows(values_only=True))
                self.assertIn("new result", rows[1])
            finally:
                result.close()
        finally:
            session.close()

    def test_locked_report_keeps_previous_file_and_progress_for_retry(self):
        source = self._xlsx("locked.xlsx")
        session = self._session()
        try:
            session.prepare(str(source))
            report = session.export_report()
            previous = report.read_bytes()
            batch_id, _total = session.begin_batch(1)
            record = session.claim_next(batch_id)
            assert record is not None
            session.commit(
                record,
                RowResult(
                    RowDisposition.RESOLVED,
                    comment="durable comment",
                    values={"portal_value": "durable result"},
                ),
            )
            session.finish_batch()
            with patch(
                "kantime_engine.reporting.os.replace",
                side_effect=PermissionError("locked"),
            ):
                with self.assertRaisesRegex(WorkbookError, "Close the report"):
                    session.export_report()
            self.assertEqual(report.read_bytes(), previous)
            self.assertEqual(session.store.completed_count(session.run.id), 1)
            retried = session.export_report()
        finally:
            session.close()

        result = load_workbook(retried, data_only=True)
        try:
            values = list(result["Results"].iter_rows(values_only=True))
            self.assertEqual(values[1][-2:], ("durable comment", "durable result"))
        finally:
            result.close()


if __name__ == "__main__":
    unittest.main()
