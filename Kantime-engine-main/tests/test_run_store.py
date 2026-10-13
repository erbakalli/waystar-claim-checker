from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tests import _bootstrap  # noqa: F401

from tests.fakes import SPEC
from kantime_engine.paths import ApplicationPaths
from kantime_engine.api import RowDisposition, RowResult
from kantime_engine.secure_data import protect, unprotect
from kantime_engine.workbook import WorkbookSession


class RunStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {"LOCALAPPDATA": self.temp.name})
        self.env.start()
        self.root = Path(self.temp.name)
        self.paths = ApplicationPaths("test", "Test App", self.root)

    def tearDown(self):
        self.env.stop()
        self.temp.cleanup()

    def _source(self, key: str = "SENSITIVE-INVOICE-8472") -> Path:
        path = self.root / "input.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Record ID", "Comments", "Portal Value"])
        sheet.append([key, None, None])
        workbook.save(path)
        workbook.close()
        return path

    def _session(self) -> WorkbookSession:
        return WorkbookSession(
            SPEC.workbook,
            SPEC.variants[0],
            plugin_id=SPEC.id,
            automation_name=SPEC.name,
            app_paths=self.paths,
        )

    def test_dpapi_round_trip_preserves_typed_scalars(self):
        value = {
            "text": "private",
            "number": Decimal("123.45"),
            "date": date(2026, 9, 27),
            "datetime": datetime(2026, 9, 27, 12, 30, tzinfo=timezone.utc),
            "blank": None,
            "boolean": True,
        }
        encrypted = protect(value)
        self.assertNotIn(b"private", encrypted)
        self.assertEqual(unprotect(encrypted), value)

    def test_sqlite_durability_pragmas_are_enabled(self):
        session = self._session()
        try:
            connection = session.store.connection
            self.assertEqual(connection.execute("PRAGMA journal_mode").fetchone()[0], "wal")
            self.assertEqual(connection.execute("PRAGMA synchronous").fetchone()[0], 2)
            self.assertEqual(connection.execute("PRAGMA foreign_keys").fetchone()[0], 1)
            self.assertEqual(connection.execute("PRAGMA secure_delete").fetchone()[0], 1)
            self.assertEqual(connection.execute("PRAGMA busy_timeout").fetchone()[0], 10_000)
        finally:
            session.close()

    def test_retry_later_is_not_selected_twice_in_one_batch(self):
        source = self._source()
        session = self._session()
        try:
            session.prepare(str(source))
            batch_id, total = session.begin_batch(10)
            self.assertEqual(total, 1)
            record = session.claim_next(batch_id)
            assert record is not None
            session.commit(
                record,
                RowResult(
                    RowDisposition.RETRY_LATER,
                    comment="partial private comment",
                    values={"portal_value": "partial private result"},
                ),
            )
            self.assertIsNone(session.claim_next(batch_id))
            self.assertEqual(session.finish_batch(), 1)
            report = session.export_report()
            exported = load_workbook(report, data_only=True)
            try:
                row = list(exported["Results"].iter_rows(values_only=True))[1]
                self.assertEqual(row[1:], ("partial private comment", "partial private result"))
            finally:
                exported.close()

            second_batch, _total = session.begin_batch(10)
            retried = session.claim_next(second_batch)
            self.assertIsNotNone(retried)
            attempts = session.store.connection.execute(
                "SELECT attempts FROM rows WHERE run_id=? AND ordinal=0",
                (session.run.id,),
            ).fetchone()[0]
            self.assertEqual(attempts, 2)
        finally:
            session.close()

    def test_recovery_releases_abandoned_in_progress_row(self):
        source = self._source()
        first = self._session()
        first.prepare(str(source))
        run_id = first.run.id
        batch_id, _total = first.begin_batch(1)
        self.assertIsNotNone(first.claim_next(batch_id))
        first.close()  # simulate a process ending after the durable claim

        resumed = self._session()
        try:
            prepared = resumed.prepare(str(source), lambda _done, _total: True)
            self.assertTrue(prepared.resumed)
            self.assertEqual(resumed.run.id, run_id)
            state = resumed.store.connection.execute(
                "SELECT status FROM rows WHERE run_id=? AND ordinal=0", (run_id,)
            ).fetchone()[0]
            self.assertEqual(state, "PENDING")
        finally:
            resumed.close()

    def test_database_files_do_not_contain_input_comment_or_result_plaintext(self):
        source = self._source()
        session = self._session()
        try:
            session.prepare(str(source))
            batch_id, _total = session.begin_batch(1)
            record = session.claim_next(batch_id)
            assert record is not None
            session.commit(
                record,
                RowResult(
                    RowDisposition.RESOLVED,
                    comment="SECRET-COMMENT-1298",
                    values={"portal_value": "SECRET-RESULT-7643"},
                ),
            )
            session.finish_batch()
        finally:
            session.close()

        database_files = list((self.root / "Test App").glob("runs.sqlite3*"))
        raw = b"".join(path.read_bytes() for path in database_files)
        for secret in (
            b"SENSITIVE-INVOICE-8472",
            b"SECRET-COMMENT-1298",
            b"SECRET-RESULT-7643",
        ):
            self.assertNotIn(secret, raw)

    def test_cleanup_removes_old_completed_snapshot_but_keeps_report(self):
        source = self._source("INV-1")
        session = self._session()
        try:
            session.prepare(str(source))
            run_id = session.run.id
            snapshot_dir = Path(session.run.snapshot_path).parent
            batch_id, _total = session.begin_batch(1)
            record = session.claim_next(batch_id)
            assert record is not None
            session.commit(record, RowResult(RowDisposition.RESOLVED))
            self.assertEqual(session.finish_batch(), 0)
            report = session.export_report()
            session.store.connection.execute(
                "UPDATE runs SET updated_at='2000-01-01T00:00:00+00:00' WHERE id=?",
                (run_id,),
            )
            session.store.connection.commit()
            self.assertEqual(session.store.cleanup(), 1)
            self.assertFalse(snapshot_dir.exists())
            self.assertTrue(report.exists())
        finally:
            session.close()


if __name__ == "__main__":
    unittest.main()
